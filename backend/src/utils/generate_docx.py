import tempfile
from datetime import datetime
from pathlib import Path
from typing import List, Dict

from docx import Document
from docx.shared import Pt, RGBColor

async def generate_docx(paragraphs: List[Dict], title: str) -> str:
    """
    Generate Word document from paragraphs
    
    Args:
        paragraphs: List of paragraph dictionaries (order, heading, heading_level, content) from Claude
        title: Document title
        
    Returns:
        str: Path to generated .docx file
    """

    doc = Document()

    # Title
    doc.add_heading(title, level = 0)

    # Content
    for paragraph in paragraphs:
        if paragraph.get("heading"):
            doc.add_heading(paragraph["heading"], level = paragraph.get("heading_level", 1))
        if paragraph.get("content"):
            doc.add_paragraph(paragraph["content"])
    
    temp_dir = Path(tempfile.gettempdir())
    filename = f"{title.replace(" ", "-")}_{datetime.now().strftime("%Y%m%d_%H%M%S")}.docx"
    filepath = str(temp_dir / filename)
    doc.save(filepath)

    return filepath
