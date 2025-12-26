import tempfile
from datetime import datetime
from pathlib import Path
from typing import List, Dict

from pptx import Presentation
from pptx.util import Pt

async def generate_pptx(slides: List[Dict], title: str) -> str:
    """
    Generate PowerPoint presentation from slide data
    
    Args:
        slides: List of slide dictionaries from Claude
        title: Presentation title
        
    Returns:
        str: Path to generated .pptx file
    """

    prs = Presentation()

    # Title Slide
    title_slide_layout = prs.slide_layouts[0]
    title_slide = prs.slides.add_slide(title_slide_layout)
    title_slide.placeholders[0].text = title
    title_slide.placeholders[1].text = datetime.now().strftime("%B %d, %Y")

    # Content Slides
    content_slide_layout = prs.slide_layouts[1]
    for slide in slides:
        content_slide = prs.slides.add_slide(content_slide_layout)
        content_slide.shapes.title.text = slide.get("title", "Untitled")

        # Bullet Points for Each Slide
        content = content_slide.placeholders[1]
        text_frame = content.text_frame
        text_frame.clear()

        for i, bullet in enumerate(slide.get("body", [])):
            if i == 0:
                p = text_frame.paragraphs[0]
            else:
                p = text_frame.add_paragraph()
        
            p.text = bullet
            p.level = 0
            p.font.size = Pt(18)
    
    # Save File
    temp_dir = Path(tempfile.gettempdir())
    filename = f"{title.replace(" ", "-")}_{datetime.now().strftime("%Y%m%d_%H%M%S")}.pptx"
    filepath = str(temp_dir / filename)
    prs.save(filepath)

    return filepath

